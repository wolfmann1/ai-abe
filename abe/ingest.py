"""Read PDF, DOCX, TXT and Markdown files and split them into chunks for retrieval."""

from __future__ import annotations

import re
from dataclasses import asdict, dataclass
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class Chunk:
    id: str
    source: str
    ordinal: int
    text: str

    def to_dict(self) -> dict:
        return asdict(self)


def read_document(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if suffix == ".docx":
        import docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n\n".join(parts)
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported file type: {path.name} (supported: {', '.join(sorted(SUPPORTED_SUFFIXES))})")


def chunk_text(text: str, size: int = 800, overlap: int = 150) -> list[str]:
    """Split text into chunks of roughly `size` characters on paragraph boundaries.

    Paragraphs longer than `size` are split on sentence boundaries, and as a last
    resort on whitespace. Adjacent chunks share up to `overlap` trailing characters
    so that an answer spanning a boundary is still retrievable.
    """
    if size <= 0:
        raise ValueError("size must be positive")
    overlap = max(0, min(overlap, size // 2))
    text = re.sub(r"\r\n?", "\n", text)
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]

    pieces: list[str] = []
    for para in paragraphs:
        para = re.sub(r"[ \t]*\n[ \t]*", " ", para)
        if len(para) <= size:
            pieces.append(para)
            continue
        sentences = re.split(r"(?<=[.!?])\s+", para)
        for sentence in sentences:
            while len(sentence) > size:
                cut = sentence.rfind(" ", 0, size)
                cut = cut if cut > 0 else size
                pieces.append(sentence[:cut].strip())
                sentence = sentence[cut:].strip()
            if sentence:
                pieces.append(sentence)

    chunks: list[str] = []
    current = ""
    for piece in pieces:
        candidate = f"{current}\n\n{piece}" if current else piece
        if len(candidate) <= size:
            current = candidate
            continue
        if current:
            chunks.append(current)
            tail = _sentence_tail(current, overlap)
            current = f"{tail} {piece}".strip() if tail and len(tail) + len(piece) + 1 <= size else piece
        else:
            current = piece
    if current:
        chunks.append(current)
    return chunks


def _sentence_tail(text: str, overlap: int) -> str:
    """The last whole sentences of `text` that fit within `overlap` characters."""
    if not overlap:
        return ""
    tail = text[-overlap:]
    boundary = re.search(r"(?<=[.!?])\s+", tail)
    return tail[boundary.end() :] if boundary else ""


def ingest_paths(paths: list[Path], size: int = 800, overlap: int = 150) -> list[Chunk]:
    chunks: list[Chunk] = []
    for path in sorted(paths):
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            continue
        text = read_document(path)
        for i, piece in enumerate(chunk_text(text, size, overlap)):
            chunks.append(Chunk(id=f"{path.name}#{i}", source=path.name, ordinal=i, text=piece))
    return chunks


def ingest_directory(directory: str | Path, size: int = 800, overlap: int = 150) -> list[Chunk]:
    directory = Path(directory)
    return ingest_paths([p for p in directory.rglob("*") if p.is_file()], size, overlap)
