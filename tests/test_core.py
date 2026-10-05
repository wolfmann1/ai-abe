import unittest

from abe.agent import Agent
from abe.cost import CostInputs
from abe.evals import EvalCase, TestSet, run_eval
from abe.index import BM25Index, tokenize
from abe.ingest import Chunk, chunk_text, read_document
from abe.intake import Criterion, Intake
from abe.providers import ExtractiveProvider
from abe.spec import AgentSpec, ProviderConfig, slugify

from .helpers import TempDir, build_example


class IngestTests(unittest.TestCase):
    def test_reads_txt_and_docx(self):
        import docx

        with TempDir() as tmp:
            (tmp / "a.txt").write_text("Plain text file.", encoding="utf-8")
            document = docx.Document()
            document.add_paragraph("Word paragraph about leases.")
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "Term"
            table.rows[0].cells[1].text = "10 years"
            document.save(tmp / "b.docx")
            self.assertEqual(read_document(tmp / "a.txt"), "Plain text file.")
            text = read_document(tmp / "b.docx")
            self.assertIn("Word paragraph about leases.", text)
            self.assertIn("Term | 10 years", text)

    def test_reads_pdf(self):
        try:
            from reportlab.pdfgen import canvas
        except ImportError:
            self.skipTest("reportlab not installed")
        with TempDir() as tmp:
            path = tmp / "c.pdf"
            pdf = canvas.Canvas(str(path))
            pdf.drawString(72, 720, "Renewal notice is due 180 days before expiry.")
            pdf.save()
            self.assertIn("180 days", read_document(path))

    def test_rejects_unknown_type(self):
        with TempDir() as tmp:
            (tmp / "x.xlsx").write_bytes(b"")
            with self.assertRaises(ValueError):
                read_document(tmp / "x.xlsx")

    def test_chunks_respect_size_and_keep_all_text(self):
        text = "\n\n".join(f"Paragraph {i}. " + "word " * 40 for i in range(20))
        chunks = chunk_text(text, size=300, overlap=60)
        self.assertTrue(all(len(c) <= 300 for c in chunks))
        joined = " ".join(chunks)
        for i in range(20):
            self.assertIn(f"Paragraph {i}.", joined)

    def test_long_sentence_is_split(self):
        chunks = chunk_text("x" * 50 + " " + "y" * 900, size=200, overlap=0)
        self.assertTrue(all(len(c) <= 200 for c in chunks))


class IndexTests(unittest.TestCase):
    def setUp(self):
        self.index = BM25Index(
            [
                Chunk("a#0", "a.txt", 0, "Restart the replication worker when a replica lags."),
                Chunk("b#0", "b.txt", 0, "Purge old artifacts when the disk is full."),
                Chunk("c#0", "c.txt", 0, "The maintenance window is on Sunday."),
            ]
        )

    def test_tokenize_drops_stopwords_keeps_dotted_terms(self):
        self.assertEqual(tokenize("What is the stash-admin v1.2 tool?"), ["stash-admin", "v1.2", "tool"])

    def test_ranks_relevant_chunk_first(self):
        hits = self.index.search("disk full purge")
        self.assertEqual(hits[0][0].source, "b.txt")

    def test_no_match_returns_empty(self):
        self.assertEqual(self.index.search("quantum chromodynamics"), [])

    def test_round_trip(self):
        with TempDir() as tmp:
            self.index.save(tmp / "i.json")
            loaded = BM25Index.load(tmp / "i.json")
            self.assertEqual(loaded.search("Sunday")[0][0].id, "c#0")


class SpecTests(unittest.TestCase):
    def test_slug_and_prompt(self):
        spec = AgentSpec(name="Lease Q&A Assistant!", out_of_scope=["Legal advice"], human_review=True)
        self.assertEqual(spec.slug, "lease-q-a-assistant")
        prompt = spec.system_prompt()
        self.assertIn("Legal advice", prompt)
        self.assertIn("Review required", prompt)
        self.assertIn("[1]", prompt)

    def test_save_load(self):
        with TempDir() as tmp:
            spec = AgentSpec(name="X", provider=ProviderConfig(kind="anthropic", model="m"))
            spec.save(tmp / "agent.yaml")
            self.assertEqual(AgentSpec.load(tmp / "agent.yaml").provider.kind, "anthropic")

    def test_slugify_fallback(self):
        self.assertEqual(slugify("!!!"), "agent")


class AgentTests(unittest.TestCase):
    def test_extractive_answer_cites_source(self):
        with TempDir() as tmp:
            agent = Agent.from_directory(build_example(tmp), ExtractiveProvider(ProviderConfig(kind="extractive")))
            answer = agent.ask("How long are release builds kept?")
            self.assertFalse(answer.declined)
            self.assertIn("365 days", answer.text)
            self.assertEqual(answer.sources[0], "faq.txt")

    def test_declines_when_nothing_matches(self):
        with TempDir() as tmp:
            agent = Agent.from_directory(build_example(tmp), ExtractiveProvider(ProviderConfig(kind="extractive")))
            self.assertTrue(agent.ask("What will the weather be on Mars next Tuesday?").declined)

    def test_llm_path_builds_numbered_context(self):
        from abe.providers import Completion, Provider

        class Recorder(Provider):
            kind = "recorder"

            def complete(self, system, user):
                self.system, self.user = system, user
                return Completion("Release builds are kept for 365 days [1].", 900, 20)

        with TempDir() as tmp:
            recorder = Recorder(ProviderConfig(kind="extractive"))
            agent = Agent.from_directory(build_example(tmp), recorder)
            answer = agent.ask("How long are release builds kept?")
            self.assertIn("[1] (source: faq.txt)", recorder.user)
            self.assertIn("Answer only from the numbered context", recorder.system)
            self.assertEqual(answer.sources, ["faq.txt"])
            self.assertEqual((answer.input_tokens, answer.output_tokens), (900, 20))


class EvalTests(unittest.TestCase):
    def test_scores_cases_and_retrieval(self):
        with TempDir() as tmp:
            agent = Agent.from_directory(build_example(tmp), ExtractiveProvider(ProviderConfig(kind="extractive")))
            testset = TestSet(
                threshold=0.5,
                cases=[
                    EvalCase(id="ok", question="How long are release builds kept?",
                             expect_contains=["365 days"], expect_source="faq.txt"),
                    EvalCase(id="miss", question="How long are release builds kept?", expect_contains=["banana"]),
                    EvalCase(id="decline", question="What will the weather be on Mars?", expect_decline=True),
                ],
            )
            report = run_eval(agent, testset)
            self.assertEqual([r.passed for r in report.results], [True, False, True])
            self.assertAlmostEqual(report.accuracy, 2 / 3)
            self.assertEqual(report.retrieval_hit_rate, 1.0)
            self.assertTrue(report.passed)
            self.assertIn("missing 'banana'", report.to_markdown())

    def test_generated_testset_loads(self):
        with TempDir() as tmp:
            testset = TestSet.load(build_example(tmp) / "evals" / "testset.yaml")
            self.assertEqual(len(testset.cases), 6)
            self.assertTrue(testset.cases[-1].expect_decline)


class IntakeAndCostTests(unittest.TestCase):
    def test_recommendations(self):
        self.assertEqual(Intake().recommendation()[0], "Time-boxed prototype")
        strong = Intake(**{k: Criterion(score=4) for k in ("data_readiness", "technical_feasibility",
                                                             "business_readiness", "minimum_accuracy",
                                                             "verifiability", "build_vs_buy")})
        self.assertEqual(strong.recommendation()[0], "Proceed")
        blocked = Intake(data_readiness=Criterion(score=1))
        self.assertEqual(blocked.recommendation()[0], "Do not build yet")
        self.assertIn("Data readiness", blocked.recommendation()[1])

    def test_cost(self):
        spec = AgentSpec(name="X")
        inputs = CostInputs.estimate_for(spec, users=10, questions_per_user_per_day=2,
                                         price_per_million_input=1.0, price_per_million_output=4.0)
        self.assertEqual(inputs.questions_per_month, 420)
        expected = 420 * inputs.avg_input_tokens / 1e6 * 1.0 + 420 * inputs.avg_output_tokens / 1e6 * 4.0
        self.assertAlmostEqual(inputs.monthly_cost, expected)
        self.assertIn("Estimated monthly model cost", inputs.to_markdown("X"))

    def test_cost_without_prices(self):
        md = CostInputs.estimate_for(AgentSpec(name="X")).to_markdown("X")
        self.assertIn("Enter your provider's per-million-token rates", md)


class BuilderTests(unittest.TestCase):
    def test_package_contents(self):
        with TempDir() as tmp:
            agent_dir = build_example(tmp)
            for name in ("agent.yaml", "blueprint.yaml", "system_prompt.md", "index.json", "intake.md",
                         "cost.md", "mcp.json", "README.md", "evals/testset.yaml"):
                self.assertTrue((agent_dir / name).exists(), name)
            self.assertEqual(len(list((agent_dir / "docs").iterdir())), 3)
            self.assertIn("Proceed", (agent_dir / "intake.md").read_text())

    def test_zip(self):
        import zipfile

        from abe.builder import zip_agent

        with TempDir() as tmp:
            archive = zip_agent(build_example(tmp))
            names = zipfile.ZipFile(archive).namelist()
            self.assertIn("stash-ops-assistant/agent.yaml", names)


if __name__ == "__main__":
    unittest.main()

